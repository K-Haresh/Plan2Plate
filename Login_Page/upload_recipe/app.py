from flask import Flask, render_template, request, redirect, url_for
from openpyxl import load_workbook

app = Flask(__name__)

# Define the Excel file path (make sure it exists or is created before running)
EXCEL_FILE_PATH = 'recipes_data_processing.xlsx'


# Function to write data to Excel
def write_to_excel(data):
    try:
        # Load the existing workbook
        wb = load_workbook(EXCEL_FILE_PATH)

        # Get the first sheet (or change this to the correct sheet name if needed)
        sheet = wb.active  # Use the active sheet, or specify sheet name like: wb['Sheet1']

        # Find the last row with data
        last_row = sheet.max_row

        # Append the data in the next available row
        sheet.append([
            data['recipeName'],
            data['ingredients'],
            data['directions'],
            data['link'],
            data['ner'],
            data['site']
        ])

        # Save the changes to the Excel file
        wb.save(EXCEL_FILE_PATH)
        print("Data appended successfully!")

    except Exception as e:
        print(f"Error saving to Excel: {e}")


@app.route("/", methods=["GET", "POST"])
def upload_recipe():
    if request.method == "POST":
        # Get the data from the form
        data = {
            'recipeName': request.form['recipeName'],
            'ingredients': request.form['ingredients'],
            'directions': request.form['directions'],
            'link': request.form['link'],
            'ner': request.form['ner'],
            'site': request.form['site']
        }

        # Write the data to the Excel sheet
        write_to_excel(data)

        # Redirect to the same page or you can show a success message
        return redirect(url_for('upload_recipe'))

    # Render the HTML form when GET request is made
    return render_template("upload_recipe.html")


if __name__ == "__main__":
    app.run(debug=True)